from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


OUTPUT_FILE = Path("dist/reuniao-docentes-ifsp-matao.pptx")

# Cores
BG = RGBColor(248, 250, 252)          # off-white
TITLE = RGBColor(15, 52, 96)          # azul petróleo
ACCENT = RGBColor(34, 139, 107)       # verde institucional
TEXT = RGBColor(51, 65, 85)           # cinza escuro
MUTED = RGBColor(100, 116, 139)       # cinza médio


def add_header_band(slide):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0), Inches(0), Inches(13.33), Inches(0.45)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()


def style_slide_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title(slide, text):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.8), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Calibri"
    p.font.bold = True
    p.font.size = Pt(34)
    p.font.color.rgb = TITLE


def add_subtitle(slide, text):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.8), Inches(0.7))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Calibri"
    p.font.size = Pt(19)
    p.font.color.rgb = MUTED


def add_bullets(slide, bullets):
    box = slide.shapes.add_textbox(Inches(1.0), Inches(2.35), Inches(11.2), Inches(4.3))
    tf = box.text_frame
    tf.clear()

    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.name = "Calibri"
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT
        p.space_after = Pt(12)


def add_footer(slide):
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(12), Inches(0.35))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = "IFSP Câmpus Matão • Engenharia de Energias Renováveis"
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.color.rgb = MUTED


def add_slide(prs, title, subtitle=None, bullets=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    style_slide_background(slide)
    add_header_band(slide)
    add_title(slide, title)
    if subtitle:
        add_subtitle(slide, subtitle)
    if bullets:
        add_bullets(slide, bullets)
    add_footer(slide)


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    add_slide(
        prs,
        "Reunião com Docentes",
        "Curso de Engenharia de Energias Renováveis – IFSP Câmpus Matão",
        ["Início da nova coordenação e alinhamento de prioridades."],
    )

    add_slide(
        prs,
        "Apresentação da coordenação",
        bullets=[
            "Fortalecer o curso e ampliar a organização interna.",
            "Construir trabalho conjunto com os docentes.",
            "Alinhar prioridades, ouvir demandas e definir encaminhamentos.",
        ],
    )

    add_slide(
        prs,
        "Visão inicial para o curso",
        bullets=[
            "Maior centralidade e consolidação no câmpus.",
            "Fortalecimento institucional do curso.",
            "Permanência e engajamento dos estudantes.",
            "Organização de demandas acadêmicas e estruturais.",
        ],
    )

    add_slide(
        prs,
        "Permanência dos estudantes",
        bullets=[
            "Prioridade para estudantes, especialmente ingressantes.",
            "Acolhimento e orientação clara desde o início.",
            "Incentivo contínuo e atenção às dificuldades acadêmicas.",
        ],
    )

    add_slide(
        prs,
        "Papel dos docentes",
        bullets=[
            "Apoio docente é fundamental para a permanência.",
            "Incentivo aos alunos e clareza nas orientações.",
            "Identificação precoce de dificuldades.",
            "Construção de ambiente acadêmico acolhedor.",
        ],
    )

    add_slide(
        prs,
        "Disciplinas práticas",
        bullets=[
            "Efetivar atividades práticas previstas nas disciplinas.",
            "Garantir a dimensão prática prevista no PPC.",
            "Integrar teoria e prática na formação do estudante.",
        ],
    )

    add_slide(
        prs,
        "Levantamento de demandas",
        bullets=[
            "Mapeamento das dificuldades nas disciplinas práticas.",
            "Equipamentos, materiais, laboratórios e apoio técnico.",
            "Necessidades de capacitação e entraves pedagógicos/operacionais.",
        ],
    )

    add_slide(
        prs,
        "Recomposição do Colegiado e do NDE",
        bullets=[
            "Recomposição do Colegiado do Curso e do NDE.",
            "Instâncias essenciais ao acompanhamento acadêmico e estratégico.",
            "Abertura para manifestação de interesse dos docentes.",
        ],
    )

    add_slide(
        prs,
        "Encaminhamentos finais",
        bullets=[
            "Envio de formulários pós-reunião.",
            "Levantamento de demandas das disciplinas práticas.",
            "Manifestação de interesse no Colegiado e NDE.",
            "Organizar, priorizar e avançar com ações concretas.",
        ],
    )

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_FILE))
    print(f"Apresentação gerada em: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_presentation()
